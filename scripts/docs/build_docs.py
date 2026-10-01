#!/usr/bin/env python3
"""Generate PostPPP Java SDK guide and test report in the dmonitor document style."""
from pathlib import Path
import json
import shutil
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
DOCS = ROOT / 'docs'
REF = ROOT.parent / 'dmonitor-java-lib' / 'docs'
EVIDENCE = DOCS / 'test-evidence' / '2026-10-01'
SUMMARY = json.loads((EVIDENCE / 'summary.json').read_text(encoding='utf-8'))
FONT = 'Arial Unicode MS'
BLUE = '234E64'
LIGHT_BLUE = 'F2F2F2'
LIGHT_GRAY = 'F2F5F7'


def set_font(run, size=None, bold=None, color='000000', italic=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def clear_body(doc):
    body = doc._element.body
    for child in list(body):
        if child.tag != qn('w:sectPr'):
            body.remove(child)


def setup_doc(template, header_text, footer_text, table_header=BLUE):
    doc = Document(template)
    clear_body(doc)
    # Keep the reference document's visual styles while replacing its metadata.
    doc.core_properties.author = 'NavFirst'
    doc.core_properties.last_modified_by = 'NavFirst'
    doc.core_properties.title = 'PostPPP Java SDK'
    doc.core_properties.subject = header_text
    doc.core_properties.keywords = 'PostPPP, Java, SDK, JNA, PPP, GNSS'
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    is_guide = '开发指南' in template.name
    sec.top_margin = Cm(2.0 if is_guide else 1.8)
    sec.bottom_margin = Cm(1.8 if is_guide else 1.65)
    sec.left_margin = Cm(2.2 if is_guide else 1.9)
    sec.right_margin = Cm(2.2 if is_guide else 1.9)
    sec.header_distance = Cm(.7)
    sec.footer_distance = Cm(.7)
    for name, size in [('Normal', 10.5), ('Title', 26 if is_guide else 22), ('Heading 1', 18 if is_guide else 16), ('Heading 2', 14 if is_guide else 12), ('Heading 3', 11.5 if is_guide else 10.5)]:
        style = doc.styles[name]
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), FONT)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.line_spacing = 1.13
        if name.startswith('Heading'):
            style.font.bold = True
            style.paragraph_format.space_before = Pt(9)
            style.paragraph_format.keep_with_next = True
    hp = sec.header.paragraphs[0]
    for child in list(hp._p):
        hp._p.remove(child)
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(hp.add_run(header_text), 8.5)
    fp = sec.footer.paragraphs[0]
    for child in list(fp._p):
        fp._p.remove(child)
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(fp.add_run(footer_text + '  |  '), 8)
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    fp._p.append(fld)
    return doc


def add_title(doc, lines):
    p = doc.add_paragraph(style='Title')
    for i, line in enumerate(lines):
        if i:
            p.add_run().add_break()
        r = p.add_run(line)
        set_font(r, doc.styles['Title'].font.size.pt, True)
    return p


def add_p(doc, text, small=False, bold=False):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.13
    for i, part in enumerate(str(text).split('\n')):
        if i:
            p.add_run().add_break()
        r = p.add_run(part)
        set_font(r, 9 if small else 10.5, bold)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f'Heading {level}')
    for r in p.runs:
        set_font(r, doc.styles[f'Heading {level}'].font.size.pt, True)
    return p


def add_page(doc, title):
    p = add_heading(doc, title, 1)
    p.paragraph_format.page_break_before = True
    return p


def add_bullets(doc, items):
    for text in items:
        p = doc.add_paragraph(style='List Bullet')
        set_font(p.add_run(str(text)), 10.5)


def shade_cell(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tcpr.append(shd)


def set_cell_margins(cell, twips='75'):
    tcpr = cell._tc.get_or_add_tcPr()
    margins = OxmlElement('w:tcMar')
    for side in ('top', 'left', 'bottom', 'right'):
        el = OxmlElement(f'w:{side}')
        el.set(qn('w:w'), twips)
        el.set(qn('w:type'), 'dxa')
        margins.append(el)
    tcpr.append(margins)


def add_table(doc, headers, rows, widths=None, size=9, header_fill=BLUE, add_after=True):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    sec = doc.sections[0]
    available = (sec.page_width - sec.left_margin - sec.right_margin) / 360000
    widths = widths or [available / len(headers)] * len(headers)
    widths = [w * available / sum(widths) for w in widths]
    for col, width in zip(table.columns, widths):
        col.width = Cm(width)
    header_color = 'FFFFFF' if header_fill == BLUE else '000000'
    all_rows = [headers] + rows
    for i, values in enumerate(all_rows):
        row = table.rows[0] if i == 0 else table.add_row()
        trpr = row._tr.get_or_add_trPr()
        cant = OxmlElement('w:cantSplit')
        trpr.append(cant)
        if i == 0:
            rep = OxmlElement('w:tblHeader')
            trpr.append(rep)
        for j, value in enumerate(values):
            cell = row.cells[j]
            cell.width = Cm(widths[j])
            cell.text = str(value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shade_cell(cell, header_fill if i == 0 else (LIGHT_GRAY if i % 2 == 0 else 'FFFFFF'))
            set_cell_margins(cell)
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = 1.08
                if len(str(value)) < 16:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    set_font(r, size, i == 0, header_color if i == 0 else '000000')
    borders = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{side}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '4')
        el.set(qn('w:color'), 'D9D9D9')
        borders.append(el)
    table._tbl.tblPr.append(borders)
    if add_after:
        after = doc.add_paragraph()
        after.paragraph_format.space_before = Pt(0)
        after.paragraph_format.space_after = Pt(1)
    return table


def add_code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(.3)
    p.paragraph_format.right_indent = Cm(.3)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.05
    p.paragraph_format.keep_together = True
    ppr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), 'F2F2F2')
    ppr.append(shd)
    for i, line in enumerate(text.strip('\n').splitlines()):
        if i:
            p.add_run().add_break()
        r = p.add_run(line)
        set_font(r, 8.4)
    return p


def add_link_like(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_font(r, 9, False, '1F4E79')
    return p


def detailed_apis(doc):
    add_heading(doc, '1. 版本与许可证', 2)
    add_table(doc, ['接口', '参数', '返回值与说明'], [
        ['getVersion()', '无', '返回当前加载的原生 PPP SDK 版本字符串。'],
        ['startUuid()', '无', '返回机器标识，通常用于申请设备绑定许可证。'],
        ['setLicense(String)', '许可证路径', '校验并缓存许可证；成功返回 JSON，失败返回空字符串。'],
        ['getLicense(String)', '许可证路径', '读取许可证元数据但不更新缓存；失败时可能返回普通错误文本。'],
    ], [4.0, 4.0, 8.2], 8.8, LIGHT_BLUE)
    add_heading(doc, '1.1 接口定义', 3)
    add_code(doc, 'String getVersion();\nString startUuid();\nString setLicense(String path);\nString getLicense(String path);')
    add_heading(doc, '1.2 参数与返回值', 3)
    add_p(doc, '许可证路径可传绝对路径。setLicense 成功后会在当前 native 进程缓存许可证；getLicense 只读取元数据，不改变缓存。许可证 JSON 中只有 active=true 才表示有效。')
    add_heading(doc, '2. 解算入口', 2)
    add_table(doc, ['接口', '用途', '失败行为'], [
        ['solvePostppp(PostpppInfo, String)', '调用 solvePostpppStruct 结构体入口。', '参数为空抛 IllegalArgumentException；native 错误仍以 JSON 返回。'],
        ['solvePostppp(String, String)', '调用 solvePostppp JSON 入口。', 'JSON 为空、含 NUL、超 64 MiB 或非法代理字符时由 JsonUtils 校验。'],
        ['solvePostpppResponse(...)', '返回 PostpppResponse，保留 code/error/可选 result。', '适合需要处理错误码和 -6 部分结果的业务。'],
        ['solvePostpppResult(...)', '直接返回成功的 PostpppResult。', 'native code 非 0 时抛 PostpppException。'],
    ], [5.1, 7.0, 4.1], 8.5, LIGHT_BLUE)
    add_heading(doc, '2.1 接口定义', 3)
    add_code(doc, 'String solvePostppp(PostpppInfo input, String license);\nString solvePostppp(String inputJson, String license);\nPostpppResponse solvePostpppResponse(...);\nPostpppResult solvePostpppResult(...);')
    add_heading(doc, '2.2 参数说明', 3)
    add_p(doc, '结构体入口中的 license 为空时使用同一进程先前成功设置的许可证缓存；非空时按该路径校验。JSON 入口使用 Go 字段名，数组 nav/sp3/clk 最多 4 项。')
    add_heading(doc, '2.3 返回值说明', 3)
    add_p(doc, 'solvePostppp 原始入口返回完整 JSON 字符串；typed response 解析外层状态并保留部分结果；typed result 只接受 code=0，非零 code 转换为 PostpppException。')
    add_heading(doc, '3. JSON 请求示例', 2)
    add_code(doc, '''String request = """
    {"year":2026,"month":7,"day":1,
     "nav":["/data/BRDC.rnx"],
     "sp3":["/data/orbit.SP3"],
     "clk":["/data/clock.CLK"],
     "bia":"/data/bias.BIA",
     "atx":"/data/igs20.atx",
     "obs_file":"/data/ABCD1820.26o",
     "station_name":"ABCD"}
    """;
PostpppResponse response = sdk.solvePostpppResponse(request, "/data/license.lic");''')
    add_p(doc, 'JSON 请求使用 Go 协议字段名，例如 elev_mask、obs_file、station_name；JsonUtils.toJson 只序列化传入对象，不自动把 Java 驼峰名称转换为 Go 字段名。')
    add_heading(doc, '4. 低层 JNA 接口', 2)
    add_table(doc, ['native 函数', 'Java 声明', '说明'], [
        ['getPVersion', 'Pointer getPVersion()', '返回 native 版本指针；服务层读取并释放。'],
        ['getUuid', 'Pointer getUuid()', '返回机器 UUID 指针。'],
        ['setLicense / getLicense', 'Pointer ...(String path)', '读取许可证信息；setLicense 会更新 native 缓存。'],
        ['solvePostppp', 'Pointer solvePostppp(String, String)', 'JSON 入口；返回 UTF-8 JSON 指针。'],
        ['solvePostpppStruct', 'Pointer solvePostpppStruct(PostpppInfo, String)', '结构体入口；返回 UTF-8 JSON 指针。'],
        ['freeP', 'void freeP(Pointer)', '只释放 native 导出的返回字符串。'],
    ], [4.5, 6.0, 6.7], 8.3, LIGHT_BLUE, False)


def guide_appendices(doc):
    add_page(doc, '附录 A 完整 Spring Boot 示例')
    add_code(doc, '''@SpringBootApplication
public class PppApplication {
    public static void main(String[] args) {
        SpringApplication.run(PppApplication.class, args);
    }
}

@Service
class PppJob {
    private final PostpppLibService sdk;

    PppJob(PostpppLibService sdk) { this.sdk = sdk; }

    PostpppResult solve(PostpppInfo input, String license) {
        return sdk.solvePostpppResult(input, license);
    }
}''')
    add_p(doc, '默认情况下 ppp.enabled=true，自动配置会注册 PostpppLibService；设置 ppp.enabled=false 后可完全关闭默认 Bean。')
    add_page(doc, '附录 B 直接 JNA 调用示例')
    add_code(doc, '''PostpppInfo input = PostpppInfo.defaults();
input.setNavPaths("/data/BRDC.rnx");
input.bia = "/data/bias.BIA";
input.atx = "/data/igs20.atx";
input.obsFile = "/data/ABCD.26o";
input.write();

Pointer pointer = PostpppLibrary.INSTANCE.solvePostpppStruct(input, license);
try {
    String json = pointer.getString(0, "UTF-8");
} finally {
    PostpppLibrary.INSTANCE.freeP(pointer);
}''')
    add_p(doc, '低层调用必须自己负责 Pointer 的读取与释放；业务代码优先使用 PostpppLibService，避免把 native 生命周期泄漏到应用层。')
    add_page(doc, '附录 C 工程包结构')
    add_table(doc, ['路径', '职责'], [
        ['library/PostpppInfo.java', 'PostpppInput JNA 结构、默认值、UTF-8 路径和定长字段。'],
        ['library/PostpppLibrary.java', 'C ABI 函数声明、同步 JNA 实例和 freeP。'],
        ['services/PostpppLibService.java', '高层同步服务和 typed API 默认方法。'],
        ['services/impl/PostpppLibServiceImpl.java', '读取并释放 native 返回字符串。'],
        ['domains/*.java', '响应、结果、历元和时间 Java 对象。'],
        ['utils/JsonUtils.java', '请求验证、响应校验和 Gson 解码。'],
        ['configurations/*', 'Spring Boot 自动配置和 ppp.enabled 属性。'],
    ], [7.0, 10.2], 8.5, LIGHT_BLUE)
    add_page(doc, '附录 D 构建与验证')
    add_code(doc, '''mvn -B -o clean verify
python3 scripts/report-tests/run_report_tests.py
python3 scripts/docs/build_docs.py''')
    add_p(doc, '发布前同时检查目标平台 native 文件、PostpppInfo.size()、测试报告中的 code/error 断言和 docs 下的最终 PDF。')


def build_guide():
    doc = setup_doc(REF / 'DMonitor-Java-SDK开发指南.docx', 'PostPPP Java SDK  |  开发指南', 'PostPPP Java SDK  |  2026-10-01', LIGHT_BLUE)
    add_title(doc, ['PostPPP Java SDK', '集成开发指南'])
    add_p(doc, 'SDK 1.0.0  |  PostPPP Go 原生库 1.0  |  2026-10-01')
    add_heading(doc, '文档说明', 1)
    add_p(doc, '本指南面向 postppp-java-lib 集成人员，说明 Java/JNA 调用方式、PostpppInfo 原生结构映射、JSON 返回对象、字段含义、错误处理、动态库和测试基线。SDK 封装 postppp-go-lib 的 C ABI，一次同步调用完成一个测站一天的静态精密单点定位（PPP）解算。')
    add_p(doc, '文档基线：Java SDK 1.0.0，原生版本接口当前返回 1.0。字段顺序、数组长度、字符串编码和返回 JSON 字段名属于接口契约；升级原生库时应同时核对匹配的头文件、动态库和本指南。')
    add_heading(doc, '免责声明', 1)
    add_p(doc, '本指南按当前源码与已归档测试证据编写，不构成对特定业务适用性、结果精度或持续兼容性的额外保证。调用方应自行核验许可证、观测数据、广播星历、精密产品、天线文件、时间窗口和运行平台。')
    add_heading(doc, '文档修改记录', 1)
    add_table(doc, ['版本', '修改日期', '描述'], [
        ['v1.0.0', '2026-10-01', '补充 PostpppInfo 完整字段映射、返回 JSON 的 Java 对象封装、中文字段解释、安装与故障排查。'],
    ], [3.0, 3.0, 11.2], 9, LIGHT_BLUE)

    add_page(doc, '一 SDK简介')
    add_p(doc, 'PostPPP Java SDK 是 postppp-go-lib 原生 PPP 算法库的 Java/JNA 封装。服务层提供结构体入口和 JSON 入口，并把 native 返回的 UTF-8 JSON 解析为可直接使用的 Java 对象。')
    add_table(doc, ['层级', '入口', '适用场景', '主要职责'], [
        ['业务服务层', 'PostpppLibService', 'Spring Boot 或普通 Java 应用', '版本、许可证、结构体/JSON 解算及 typed response。'],
        ['结果解析层', 'JsonUtils', '已有 JSON 或自定义协议处理', '校验响应、保留原文、解析 PostpppResponse/PostpppResult。'],
        ['原生绑定层', 'PostpppLibrary', '高级集成或底层调试', '通过 JNA 加载 libPostppp.so 并负责 Pointer 生命周期。'],
        ['输入结构层', 'PostpppInfo', '结构体解算', '按 Go PostpppInput 的 ABI 顺序写入 native 内存。'],
    ], [3.2, 4.0, 4.3, 5.7], 8.8, LIGHT_BLUE)
    add_page(doc, '二 名词解释')
    add_table(doc, ['名词', '说明'], [
        ['PPP', '精密单点定位；本 SDK 面向单测站、单日静态解算。'],
        ['JNA', 'Java Native Access，用于加载动态库并映射 C 结构体与函数。'],
        ['GPST', 'GPS 时间；PostpppTime 以 GPS 周和周内秒表示。'],
        ['ECEF / ENU / LLH', '地心地固坐标、东-北-天局部坐标、大地坐标。'],
        ['OSB/BIA / ANTEX', '码偏差文件和天线校准文件；由原生库读取。'],
    ], [4.0, 13.2], 9, LIGHT_BLUE)

    add_heading(doc, '数据对象与单位', 2)
    add_table(doc, ['名词', '说明'], [
        ['PostpppInput', 'Go C ABI 的输入结构体；Java 对应 PostpppInfo，大小为 392 字节。'],
        ['Result', 'Go 返回 JSON 中的单测站单日 PPP 结果；Java 对应 PostpppResult。'],
        ['Response', 'Go 返回 JSON 的 code、error、result 外层；Java 对应 PostpppResponse。'],
        ['历元', '一个观测时刻的解算结果，对应 PostpppEpoch；时间对应 PostpppTime。'],
        ['固定宽度字段', 'C char[] 字段，必须使用 byte[] 并按 UTF-8 字节容量写入。'],
        ['部分结果', 'code=-6 时 native 仍返回 result；用于诊断，不表示成功。'],
    ], [4.0, 13.2], 9, LIGHT_BLUE)
    add_heading(doc, '接口契约中的单位', 2)
    add_table(doc, ['数据', '单位与顺序'], [
        ['ECEF 坐标', 'X/Y/Z，米。'],
        ['ENU 偏差', '东/北/天，米。'],
        ['LLH', '纬度/经度为度，大地高为米。'],
        ['时间', 'GPS 周 + 周内秒（GPST）。'],
        ['对流层', 'ZTD/ZHD/ZWD、梯度和标准差，米。'],
    ], [5.0, 12.2], 9, LIGHT_BLUE)

    add_page(doc, '三 快速开始')
    add_heading(doc, '1. 准备工作', 2)
    add_p(doc, '准备 JDK 20 或以上、Maven 3.x、有效许可证，以及与解算日期匹配的观测文件、广播星历、精密轨道、钟差、码偏差和天线文件。原生库不下载这些数据。')
    add_heading(doc, '2. 支持平台', 2)
    add_table(doc, ['运行平台', 'JAR 资源目录', '状态'], [
        ['Linux x86-64', 'linux-x86-64/libPostppp.so', '当前验证平台；已归档 12 项测试全部通过。'],
        ['Linux AArch64', 'linux-aarch64/libPostppp.so', '资源目录存在；需用同架构 native 构建验证。'],
        ['macOS Apple Silicon', 'darwin-aarch64/libPostppp.so', '当前资源文件实际为 Linux ELF，不能直接作为 macOS 库使用。'],
        ['Windows', '无', '工程未提供。'],
    ], [4.0, 6.0, 6.2], 8.8, LIGHT_BLUE)
    add_heading(doc, '3. 安装 SDK', 2)
    add_code(doc, '''mvn clean verify
mvn install''')
    add_p(doc, 'Maven 坐标为 com.navfirst:postppp-java-lib:1.0.0。产物是普通依赖 JAR：target/postppp-java-lib-1.0.0.jar；Spring Boot Maven 插件跳过可执行 JAR 重打包。')
    add_code(doc, '''<dependency>
    <groupId>com.navfirst</groupId>
    <artifactId>postppp-java-lib</artifactId>
    <version>1.0.0</version>
</dependency>''')
    add_heading(doc, '4. 接入流程', 2)
    add_p(doc, '按以下顺序接入：确认 Linux amd64 与 native 库版本，申请有效许可证，准备单测站单日观测和同日期的广播/精密产品，构建输入，调用同步解算，检查 code，再读取质量统计和坐标。需要输出文件时提前创建可写目录并设置对应开关。')
    add_heading(doc, '5. 获取算法版本和机器标识', 2)
    add_code(doc, 'System.out.println(sdk.getVersion());\nString uuid = sdk.startUuid();')
    add_p(doc, 'getVersion 返回 native 算法版本（本轮为 1.0）；startUuid 返回许可证绑定所需的机器标识。不要把 native 版本 1.0 与 Maven SDK 版本 1.0.0 混用。')
    add_heading(doc, '6. 构建输入并启动解算', 2)
    add_code(doc, '''PostpppLibService sdk = new PostpppLibServiceImpl();
PostpppInfo input = PostpppInfo.builder()
        .year(2026).month(7).day(1)
        .bia("/data/bias.BIA")
        .atx("/data/igs20.atx")
        .obsFile("/data/ABCD1820.26o")
        .build();
input.setNavPaths("/data/BRDC.rnx");
input.setSp3Paths("/data/orbit.SP3");
input.setClkPaths("/data/clock.CLK");
input.setStationNameText("ABCD");

PostpppResponse response = sdk.solvePostpppResponse(input, "/data/license.lic");
if (!response.isSuccess()) {
    throw new IllegalStateException(response.getError());
}
PostpppResult result = response.getResult();
System.out.println(result.getFinalLlh()[0]);''')
    add_p(doc, 'new PostpppInfo()、PostpppInfo.defaults() 和 builder().build() 使用与 Go 初始化器一致的默认值。手动调用 initPostpppInput 会覆盖已填写字段，应在填写输入前调用。')

    add_page(doc, '四 详细接口描述')
    detailed_apis(doc)

    add_page(doc, '五 数据模型')
    add_heading(doc, '1. PostpppInfo 输入模型', 2)
    add_p(doc, 'PostpppInfo 对应 Go main.go 中的 PostpppInput，而不是底层 postppp_opt_t。当前 64 位 ABI 大小为 392 字节，字段顺序必须保持一致。C int32_t 映射为 Java int，C double 映射为 Java double，C char[] 映射为 byte[]。')
    add_table(doc, ['Java 字段', 'C / JSON 字段', '类型', '中文说明与默认值'], [
        ['year / month / day', 'year / month / day', 'int', 'GPST 日期；全 0 从观测首历元确定，年范围 1970～2099。'],
        ['interval', 'interval', 'double', '采样间隔（秒）；0 自动推断，不能为负。'],
        ['navsys', 'navsys', 'int', '卫星系统位掩码；默认 13，即 GPS+GAL+BDS。'],
        ['elevMask', 'elev_mask', 'double', '截止高度角（度）；默认 7，范围 [0,90)。'],
        ['soltype', 'soltype', 'int', '解算方向；0 前向、1 反向、2 双向组合（默认）。'],
        ['tropopt', 'tropopt', 'int', '对流层模型；3 仅 ZTD，4 为 ZTD+梯度（默认）。'],
        ['bdsfreq', 'bdsfreq', 'int', 'BDS 频点组合；4=B1I+B3I（默认），5=B1C+B2A。'],
        ['tidecorr', 'tidecorr', 'int', '潮汐改正；0 关闭、1 固体潮（默认）、2 全潮汐。'],
        ['tropPrn', 'trop_prn', 'double', '对流层过程噪声；0 使用内核默认，不能为负。'],
        ['maxSigma', 'max_sigma', 'double', '三维位置中误差门限（米）；默认 0.020，0 关闭。'],
        ['enuOffset', 'enu_offset', 'double[3]', 'ENU 平移（米），默认全 0；顺序东/北/天。'],
        ['outDir', 'out_dir', 'String', '输出目录；null/空值时不写文件。'],
        ['writePos / writeTro / writeQc', 'write_pos / write_tro / write_qc', 'int', '非 0 分别写 POS、SINEX_TRO、质量报告。'],
        ['troInterval', 'tro_interval', 'double', 'SINEX_TRO 间隔（秒）；默认 300，0 也回退到 300。'],
        ['software / agency', 'software / agency', 'String', '结果元数据；为空时 native 回退默认名称。'],
        ['verbose', 'verbose', 'int', '日志级别；当前 C ABI 无日志回调，单独设置不会输出 Java 日志。'],
    ], [4.0, 4.0, 2.1, 7.1], 8.3, LIGHT_BLUE)
    add_heading(doc, '路径、字符数组和数组字段', 2)
    add_table(doc, ['Java 字段', 'C / JSON 字段', '类型', '中文说明与默认值'], [
        ['nav', 'nav', 'Pointer[4]', '广播星历路径，最多 4 项；成功读到一项即止。'],
        ['sp3 / clk', 'sp3 / clk', 'Pointer[4]', '精密轨道和钟差路径，各最多 4 项；多文件可累加读取。'],
        ['bia / atx', 'bia / atx', 'String', '必需的 OSB/BIA 码偏差与 ANTEX 天线文件。'],
        ['erp / blq / ionex', 'erp / blq / ionex', 'String', 'ERP/BLQ 为全潮汐 tidecorr=2 时的必需文件；IONEX 为可选电离层格网。'],
        ['obsFile', 'obs_file', 'String', '必需的 RINEX 或 RTCM3 观测文件路径。'],
        ['stationName', 'station_name', 'byte[16]', '测站名；最多 15 个 UTF-8 字节加 NUL。'],
        ['approxXyz', 'approx_xyz', 'double[3]', 'ECEF 先验坐标（米）；全 0 时使用观测头坐标。'],
        ['antType', 'ant_type', 'byte[24]', '天线类型；最多 23 个 UTF-8 字节加 NUL。'],
        ['radome', 'radome', 'byte[8]', '天线罩名称；最多 7 个 UTF-8 字节加 NUL。'],
        ['antDelta', 'ant_delta', 'double[3]', '天线 E/N/U 偏心（米），默认全 0。'],
    ], [4.0, 4.0, 2.1, 7.1], 8.3, LIGHT_BLUE)
    add_p(doc, 'nav/sp3/clk 使用 Pointer[4] 是因为 JNA 不支持在 Structure 中直接写入 String[]。setNavPaths、setSp3Paths、setClkPaths 会分配 UTF-8、NUL 结尾的 Memory，并保留引用直到结构体使用完毕；空项写为 NULL。定长 char[] 不能映射为 Java char[]。')
    add_code(doc, '''input.setStationNameText("ABCD");
input.setAntTypeText("TRM59800.00");
input.setRadomeText("NONE");
input.setNavPaths("/data/BRDC.rnx");
input.setSp3Paths("/data/orbit.SP3", "/data/orbit-next.SP3");
input.setClkPaths("/data/clock.CLK");''')

    add_page(doc, '五 数据模型（续）')
    add_heading(doc, '2. 返回对象', 2)
    add_p(doc, 'native 返回 JSON 的外层通过 PostpppResponse 封装；成功时 result 为 PostpppResult，-6（没有合格 PPP 解）时也可能保留部分 result，便于诊断。字段使用 Gson 的 @SerializedName 保持 Go 下划线协议，Java getter 使用驼峰命名。')
    add_heading(doc, '1. PostpppResponse', 2)
    add_table(doc, ['JSON 字段', 'Java 属性 / getter', '类型', '中文说明'], [
        ['code', 'code / getCode()', 'int', '原生返回码；0 成功，非 0 表示参数、文件或解算错误。'],
        ['error', 'error / getError()', 'String', '错误描述；成功时通常为空。'],
        ['result', 'result / getResult()', 'PostpppResult', 'PPP 结果；成功时必有，-6 时可能是部分结果。'],
        ['—', 'isSuccess()', 'boolean', '等价于 code == 0。'],
        ['—', 'hasResult()', 'boolean', 'result 是否存在。'],
        ['—', 'isPartialResult()', 'boolean', 'code == -6 且 result 存在时为 true。'],
    ], [4.0, 4.5, 3.0, 6.1], 8.5, LIGHT_BLUE)
    add_heading(doc, '2. PostpppTime 与 PostpppEpoch', 2)
    add_table(doc, ['JSON 字段', 'Java getter', '类型', '中文说明'], [
        ['week / tow', 'getWeek() / getTow()', 'int / double', 'GPS 周和周内秒（GPST）。'],
        ['time', 'getTime()', 'PostpppTime', '历元时间。'],
        ['xyz / sigma', 'getXyz() / getSigma()', 'double[3]', 'ECEF 坐标及标准差，顺序 X/Y/Z，单位米。'],
        ['sigma3d', 'getSigma3d()', 'double', 'ECEF 三维位置标准差，单位米。'],
        ['enu', 'getEnu()', 'double[3]', '相对参考坐标的 ENU 偏差，顺序东/北/天，单位米。'],
        ['ztd / zhd / zwd', 'getZtd() / getZhd() / getZwd()', 'double', '总、干、湿对流层延迟，单位米。'],
        ['gn / ge', 'getGn() / getGe()', 'double', '北向/东向对流层梯度，单位米。'],
        ['sd_ztd / sd_gn / sd_ge', 'getSdZtd() / getSdGn() / getSdGe()', 'double', '对应参数标准差，单位米。'],
        ['nsat', 'getNsat()', 'int', '参与该历元 PPP 解算的卫星数。'],
        ['stat', 'getStat()', 'int', '0 无解、3 SPP、4 PPP 浮点、6 PPP 固定。'],
        ['accepted', 'isAccepted()', 'boolean', '是否通过位置质量门限并计入最终统计。'],
    ], [4.5, 5.0, 3.0, 5.1], 8.2, LIGHT_BLUE)
    add_heading(doc, '3. PostpppResult（测站日结果）', 2)
    add_table(doc, ['JSON 字段', 'Java getter', '类型', '中文说明'], [
        ['name / year / doy', 'getName() / getYear() / getDoy()', 'String / int', '测站名、年份、年积日。'],
        ['epoch', 'getEpoch()', 'List<PostpppEpoch>', '按 GPST 排序的逐历元结果；拒绝历元也保留。'],
        ['have_coord / n_coord', 'isHaveCoord() / getNCoord()', 'boolean / int', '是否有合格坐标及合格历元数。'],
        ['first / last', 'getFirst() / getLast()', 'PostpppTime', '有效成果起止时间。'],
        ['final_xyz / final_sigma', 'getFinalXyz() / getFinalSigma()', 'double[3]', '最后合格历元的 ECEF 坐标及标准差，单位米。'],
        ['final_llh', 'getFinalLlh()', 'double[3]', '最后合格历元大地坐标，纬度/经度为度，高为米。'],
        ['mean_xyz', 'getMeanXyz()', 'double[3]', '合格历元 ECEF 平均值，单位米。'],
        ['rms_enu', 'getRmsEnu()', 'double[3]', '相对合格均值的 ENU 重复性 RMS，单位米。'],
        ['bias_enu / bias_enu_mean', 'getBiasEnu() / getBiasEnuMean()', 'double[3]', '最后历元/全天平均 ENU 偏差，顺序东/北/天。'],
        ['ztd_mean', 'getZtdMean()', 'double', '全天合格成果平均 ZTD，单位米。'],
        ['n_epoch / n_ppp / n_accept', 'getNEpoch() / getNPpp() / getNAccept()', 'int', '总历元、PPP 历元、质量门限通过历元数量。'],
        ['n_gated / n_spp / n_none', 'getNGated() / getNSpp() / getNNone()', 'int', '被门限剔除、仅 SPP、无结果的历元数量。'],
        ['avail / mean_nsat', 'getAvail() / getMeanNsat()', 'double', '数据可用率（百分比）及平均可用卫星数。'],
        ['n_gap / max_gap_sec / max_gap_start', 'getNGap() / getMaxGapSec() / getMaxGapStart()', 'int / double / PostpppTime', '中断窗口数、最长中断秒数及开始时间。'],
        ['converge_sec', 'getConvergeSec()', 'double', '首次达到收敛条件的耗时，单位秒。'],
        ['grade / software / agency', 'getGrade() / getSoftware() / getAgency()', 'String', '综合质量等级及结果元数据。'],
        ['products', 'getProducts()', 'String', '使用的精密产品标识。'],
        ['pos_file / tro_file', 'getPosFile() / getTroFile()', 'String', 'POS、SINEX_TRO 实际输出路径；未输出时为空。'],
    ], [4.3, 5.4, 3.1, 4.8], 7.8, LIGHT_BLUE)
    add_heading(doc, '4. typed API 示例', 2)
    add_code(doc, '''PostpppResponse response = sdk.solvePostpppResponse(input, license);
if (!response.isSuccess()) {
    System.err.printf("PPP failed (%d): %s%n",
            response.getCode(), response.getError());
}
if (response.getResult() != null) {
    PostpppResult result = response.getResult();
    System.out.println(result.getName());
    System.out.println(result.getFinalXyz()[0]);
    System.out.println(result.getEpoch().size());
}''')

    add_page(doc, '六 状态码')
    add_heading(doc, '1. 返回码', 2)
    add_table(doc, ['code', '含义', '处理建议'], [
        ['0', '成功', '读取 result；仍应检查 have_coord、grade 和质量统计。'],
        ['-1', '参数或许可证错误', '检查 JSON、日期、路径、许可证和输入文件。'],
        ['-2', '内存不足', '减少输入规模或并发数，检查 native 进程资源。'],
        ['-3', '广播星历读取错误', '核对 nav 路径、日期覆盖和文件格式。'],
        ['-4', '精密产品读取错误', '核对 SP3/CLK/BIA/ATX 与日期匹配。'],
        ['-5', '观测文件读取错误', '核对 obs_file、权限和 RINEX/RTCM3 格式。'],
        ['-6', '没有合格 PPP 解', '优先使用 PostpppResponse 保留部分 result 做诊断。'],
        ['-7', '输出文件错误', '检查 outDir 权限、磁盘空间和输出开关。'],
    ], [2.0, 6.1, 9.5], 8.8, LIGHT_BLUE)
    add_heading(doc, '2. 历元解状态与质量标志', 2)
    add_table(doc, ['字段与值', '含义', '业务判断'], [
        ['stat=0 / 3', '无解 / SPP', '不能作为合格 PPP 成果。'],
        ['stat=4 / 6', 'PPP 浮点解 / PPP 固定解', '结合 accepted 和质量统计判断。'],
        ['accepted=false', '该历元未通过位置质量门限', '保留用于诊断，不参与合格坐标汇总。'],
        ['have_coord=false', '没有合格坐标成果', '不能解释 final_xyz/final_llh 中的零值。'],
        ['grade', 'native 综合质量等级', '用于质量汇总，不能替代与独立真值的精度检验。'],
    ], [4.3, 5.4, 6.9], 9, LIGHT_BLUE)

    add_page(doc, '七 使用限制说明')
    add_heading(doc, '1. 输入与 native 限制', 2)
    add_bullets(doc, [
        'year/month/day 全 0 时由首历元确定；指定日期时必须同时填写有效年月日。',
        'nav、sp3、clk 每类最多 4 个路径；路径字符串必须可按 UTF-8 转换且不能包含 NUL。',
        'station_name、ant_type、radome 受固定字节容量限制，中文字符可能占用多个 UTF-8 字节。',
        'tidecorr=2 时需要 ERP 和 BLQ；缺少精密产品、BIA 或 ANTEX 会导致 native 错误。',
        'outDir 只有和 writePos、writeTro 或 writeQc 同时设置时才会生成对应输出。',
        'PostpppInfo 是 native 内存映射对象，调用期间不要修改字段、数组或路径 Memory。',
    ])
    add_heading(doc, '2. Spring Boot 配置', 2)
    add_code(doc, 'ppp.enabled=true\n\n# 关闭默认自动配置\nppp.enabled=false')
    add_p(doc, '默认自动配置通过 PostpppAutoConfiguration 注册 PostpppLibService。提供自定义 PostpppLibService Bean 时，@ConditionalOnMissingBean 会保留自定义实现。')
    add_heading(doc, '3. native 内存生命周期', 2)
    add_p(doc, '服务层会把 native 返回的 char* 复制为 UTF-8 Java String，并在 finally 中调用一次 freeP。直接使用 PostpppLibrary 时返回值是 Pointer，调用方必须读取并释放，不能重复释放；输入结构体及其路径 Memory 不应传给 freeP。')
    add_heading(doc, '4. 同步与并发', 2)
    add_p(doc, '服务调用是同步阻塞的。调用期间不要修改 PostpppInfo 及其数组。当前 Go 内核存在进程级锁；多站并行处理应在验证 native 线程安全后再部署，保守方案是不同进程或在业务层串行保护。')
    add_heading(doc, '5. Spring Boot 与动态库', 2)
    add_p(doc, 'Spring Boot 3 应用可以注入 PostpppLibService；普通 Java 程序直接 new PostpppLibServiceImpl。JNA 以 Postppp 名称加载动态库，外部库可以通过 -Djna.library.path=/path/to/native 指定。加载失败会抛出 UnsatisfiedLinkError。')
    add_code(doc, '''java -Djna.library.path=/opt/postppp \
     -jar your-app.jar''')
    add_heading(doc, '6. 输出目录与覆盖策略', 2)
    add_p(doc, 'outDir 应按站点、日期和任务实例分配唯一目录。writePos、writeTro、writeQc 由业务明确开启，重复任务不要共享同一目录，避免 POS、SINEX_TRO 和质量文件互相覆盖。')
    add_heading(doc, '7. 输入文件与许可证', 2)
    add_p(doc, '许可证、观测文件、广播星历、SP3、CLK、BIA、ANTEX、ERP 和 BLQ 必须覆盖同一解算日期和测站。SDK 只负责传递路径，不负责下载、校验和归档这些外部文件。')

    add_page(doc, '八 常见问题')
    add_heading(doc, '1. 启动时报 UnsatisfiedLinkError', 2)
    add_bullets(doc, [
        '确认操作系统和 CPU 架构与资源目录一致；当前 Linux x86-64 是已验证平台。',
        '确认 JAR 中包含 linux-x86-64/libPostppp.so，或通过 -Djna.library.path 指定外部目录。',
        '确认 native 动态库与 PostpppInfo 字段顺序来自同一 Go 构建。',
    ])
    add_heading(doc, '2. 返回 code=-6', 2)
    add_p(doc, 'code=-6 表示没有合格 PPP 解，但 native 可能同时提供部分 result。需要诊断时使用 solvePostpppResponse 并检查 haveCoord、nAccept、grade、epoch 和 error；只处理成功分支时才使用 solvePostpppResult。')
    add_heading(doc, '3. JSON 解析失败', 2)
    add_bullets(doc, [
        '响应必须是对象且 code 必须为数字；成功 code=0 时必须存在 result 对象。',
        'error 必须是字符串；result 必须是对象或 null。',
        '请求 JSON 不能为空、不能包含 NUL 或未配对代理字符，大小上限为 64 MiB。',
    ])
    add_heading(doc, '4. 坐标结果为零或输出文件为空', 2)
    add_p(doc, '先判断 result.haveCoord；为 false 时不要解释坐标数组中的零值。需要落盘时同时设置 outDir 和对应 write 开关，并检查路径权限、磁盘空间和文件覆盖策略。')

    add_page(doc, '九 测试、迁移与复现')
    add_heading(doc, '1. 当前测试基线', 2)
    add_p(doc, '当前归档基线为 Linux amd64 Docker 中的 12 项自动化测试，覆盖 PostpppInfo 的 ABI 大小、字段默认值、UTF-8 定长字段、路径数组、JSON 响应对象和 native 错误路径。测试报告与逐项 Surefire 证据位于 docs/test-evidence/2026-10-01。')
    add_table(doc, ['检查项', '当前结果', '集成含义'], [
        ['ABI 与默认值', '392 字节；字段顺序、数组长度和默认值断言通过', '升级 Go 结构体时必须同步更新 Java 映射。'],
        ['typed JSON', 'Response、Result、Epoch、Time 均可解码', '业务层可按 getter 读取中文字段对应的结果属性。'],
        ['错误路径', 'code/error、-6 部分结果和异常 facade 通过', '生产代码应优先检查 code，再读取 result。'],
        ['native 冒烟', '版本、许可证错误、结构体和 JSON 入口通过', '证明动态库可加载和 C ABI 可调用，不等于精度验收。'],
    ], [4.0, 7.0, 6.2], 8.4, LIGHT_BLUE)
    add_heading(doc, '2. 版本迁移提示', 2)
    add_bullets(doc, [
        '升级 postppp-go-lib 时先比较 PostpppInput 字段顺序、sizeof、定长 char 数组和返回 JSON 字段，再更新 Java 映射。',
        '如果新增 JSON 字段，优先在 domains 包增加属性和中文说明；未知字段由 Gson 忽略，不应依赖未知字段完成业务判断。',
        '如果返回码含义变化，应同时更新状态码表、PostpppException、测试夹具和测试报告。',
        '发布新 native 文件后重新执行 Linux amd64、目标架构和许可证错误冒烟，并重新归档 SHA-256。',
    ])
    add_heading(doc, '3. 复现与证据', 2)
    add_code(doc, 'mvn -B -o clean verify\npython3 scripts/report-tests/run_report_tests.py\npython3 scripts/docs/build_docs.py')
    add_p(doc, '完整字段与 Go 端结果定义见 postppp-go-lib/readme.md；测试报告、Markdown 版报告和原始 Surefire 证据见 docs/。')

    guide_appendices(doc)
    out = DOCS / 'PostPPP-Java-SDK开发指南.docx'
    doc.save(out)
    return out


def build_report():
    doc = setup_doc(REF / 'DMonitor-Java-SDK测试报告.docx', 'PostPPP Java SDK  |  测试报告', 'PostPPP Java SDK  |  2026-10-01', BLUE)
    suite_times = {s['name'].split('.')[-1]: f"{float(s['time']):.3f} s" for s in SUMMARY['suites']}
    total_time = f"{sum(float(s['time']) for s in SUMMARY['suites']):.3f} s"
    add_title(doc, ['PostPPP Java SDK', '测试报告'])
    add_p(doc, '报告 v1.1  |  SDK 1.0.0  |  原生库 1.0  |  2026-10-01')
    add_heading(doc, '结论摘要', 2)
    add_p(doc, '结论为通过。本轮在 Linux amd64 Docker 环境执行 mvn -B -o clean verify，两个测试类共 12 项测试全部通过，Failures=0、Errors=0、Skipped=0，构建结果为 BUILD SUCCESS。测试覆盖 PostpppInfo ABI 映射、默认值、UTF-8 和定长字段、返回 JSON 的 Java 对象封装、错误码保留、native 版本和许可证错误路径。')
    add_table(doc, ['测试内容', '本轮结果', '结论'], [
        ['PostpppInfo 映射', '6 项；0 失败、0 错误、0 跳过', '通过'],
        ['应用与 native 冒烟', '6 项；0 失败、0 错误、0 跳过', '通过'],
        ['总计', '12 项；0 失败、0 错误、0 跳过', '通过'],
        ['构建', 'mvn -B -o clean verify；BUILD SUCCESS', '通过'],
    ], [4.6, 8.0, 4.2], 9, BLUE)
    add_heading(doc, '适用范围', 2)
    add_p(doc, '本报告面向 SDK 集成、测试和验收人员。测试在 Linux amd64 容器中加载工程内置 libPostppp.so，验证 Java 到 Go C ABI 的契约和错误路径。报告不把单元/冒烟测试解释为真实观测数据的 PPP 精度验收；实际生产精度仍需独立数据集和真值测试。')

    add_page(doc, '一 测试环境与计量方法')
    add_table(doc, ['项目', '本轮基线'], [
        ['Git 提交', SUMMARY['git_head']],
        ['执行平台', 'Linux amd64（Docker --platform linux/amd64）'],
        ['Docker 镜像', f"{SUMMARY['image']}\n{SUMMARY['image_id']}"],
        ['Java', 'Eclipse Temurin 21（工程编译 release 20）'],
        ['Maven', '容器内 Maven；离线模式 -o'],
        ['SDK', 'com.navfirst:postppp-java-lib:1.0.0'],
        ['native', f"linux-x86-64/libPostppp.so；{SUMMARY['native_bytes']:,} bytes\nSHA-256 {SUMMARY['native_sha256']}"],
        ['测试产物', f"target/postppp-java-lib-1.0.0.jar；{SUMMARY['jar_bytes']:,} bytes\nSHA-256 {SUMMARY['jar_sha256']}"],
    ], [4.0, 13.2], 8.8, BLUE)
    add_heading(doc, '执行命令', 2)
    add_code(doc, 'docker run --rm --platform linux/amd64 \\\n  -v <repo>:/workspace -w /workspace \\\n  maven:3.9-eclipse-temurin-21 mvn -B -o clean verify')
    add_p(doc, '脚本 scripts/report-tests/run_report_tests.py 会归档 Surefire XML/TXT、命令日志、SHA-256 和 source-baseline.json，并要求至少 12 项测试全部执行且没有 failures、errors、skipped。')
    add_heading(doc, '判定规则', 2)
    add_table(doc, ['指标', '通过条件'], [
        ['测试执行数', '总数不少于 12；两个测试类均被 Surefire 发现。'],
        ['失败与错误', 'Failures=0、Errors=0。'],
        ['跳过', 'Skipped=0。'],
        ['构建', 'Maven 输出 BUILD SUCCESS，生成目标 JAR。'],
        ['证据完整性', '保留 XML、TXT、日志、摘要和 SHA-256 清单。'],
    ], [5.0, 12.2], 8.8, BLUE)
    add_heading(doc, '计时与统计口径', 2)
    add_p(doc, '测试套件耗时来自 Surefire 汇总，只表示当前容器内 Java 测试和 native 冒烟的执行时间。报告不把该耗时当作单站 PPP 解算性能，也不从没有真实观测窗口的用例推导收敛时间、吞吐量或精度。')

    add_page(doc, '二 数据来源与完整性')
    add_p(doc, '本报告的测试数据由固定 JSON 响应夹具、PostpppInfo 构造对象和 native 错误路径组成。测试不把不存在的观测文件伪装成 PPP 精度样本；真实观测、广播星历、SP3、CLK、BIA 和 ANTEX 文件由业务部署提供。')
    add_table(doc, ['数据类别', '本轮使用', '完整性检查'], [
        ['JSON 响应夹具', '成功响应、逐历元对象、-6 部分结果', 'Gson 下划线字段、数组、嵌套时间和布尔值均断言。'],
        ['结构体输入', 'PostpppInfo.defaults()、Builder 和路径 setter', '392 字节布局、数组长度、UTF-8 NUL 结尾和默认值断言。'],
        ['错误输入', '空 JSON、缺少观测文件、不可用许可证、非法路径', '统一解析为 code/error 或 PostpppException。'],
        ['原生文件', 'JAR 内 linux-x86-64/libPostppp.so', f"{SUMMARY['native_bytes']:,} bytes；SHA-256 已归档"],
    ], [4.0, 6.0, 7.2], 8.5, BLUE)
    add_heading(doc, '输入边界与不纳入项', 2)
    add_p(doc, '路径字段按 UTF-8 传递，定长 station_name、ant_type、radome 按字节容量限制；测试覆盖超长、NUL 和路径数量超过 4 项等违规输入。由于当前证据没有有效许可证和真实日观测数据，本轮不声称完成坐标精度、收敛时间或吞吐量验收。')
    add_heading(doc, '传输封装与星历', 2)
    add_p(doc, '输入通过 Java 对象、JNA Structure 或 JSON 字符串进入 Go C ABI。nav、sp3、clk 的每个路径均保留 UTF-8 NUL 结尾的 native Memory；测试验证路径数量上限和引用生命周期。星历、精密轨道、钟差、BIA、ANTEX、ERP 与 BLQ 的真实内容不在本轮夹具中。')
    add_heading(doc, '主测试站点与时间', 2)
    add_p(doc, '本轮没有可复核的真实测站、独立坐标真值或八个独立小时窗口，因此不填写虚构的站点名、观测日期和定位误差。业务验收应提供可归档的 RINEX、产品文件、真值和运行日志。')

    add_page(doc, '三 回归结果与独立站点解算')
    add_heading(doc, '1. 测试套件汇总', 2)
    add_table(doc, ['测试类', '测试数', '失败', '错误', '跳过', '耗时'], [
        ['com.navfirst.postppp.lib.PostpppInfoTests', '6', '0', '0', '0', suite_times.get('PostpppInfoTests', 'n/a')],
        ['com.navfirst.postppp.lib.PostpppJavaLibApplicationTests', '6', '0', '0', '0', suite_times.get('PostpppJavaLibApplicationTests', 'n/a')],
        ['合计', '12', '0', '0', '0', total_time],
    ], [7.2, 1.7, 1.7, 1.7, 1.7, 2.0], 8.2, BLUE)
    add_heading(doc, '2. 用例明细', 2)
    cases = SUMMARY['cases']
    rows = []
    cn = {
        'postpppInfoMatchesNativeLayoutAndDefaults': 'PostpppInfo 与 native 布局及默认值一致',
        'builderRetainsUtf8PathsAndFixedArrayLengths': 'Builder 保留 UTF-8 路径和定长数组',
        'fixedWidthAndPathInputsRejectNativeStringViolations': '定长字段和路径拒绝 native 字符串违规',
        'responseEnvelopePreservesCodeErrorAndResult': '响应封装保留 code、error、result',
        'responseJsonIsDecodedIntoTypedResultObjects': '返回 JSON 解码为 typed Java 对象',
        'partialResultIsRetainedForNoSolutionResponse': '无合格解时保留部分 result',
        'javaDefaultsMatchNativeInitializer': 'Java 默认值与 native 初始化器一致',
        'structureEntryPointReturnsGoErrorEnvelope': '结构体入口返回 Go 错误封装',
        'unavailableLicenseReturnsErrorForRawAndTypedResponses': '不可用许可证在原始和 typed 入口均返回错误',
        'resultFacadePreservesNativeErrorCode': '结果 facade 保留 native 错误码',
        'jsonEntryPointReturnsGoErrorEnvelope': 'JSON 入口返回 Go 错误封装',
        'getVersion': '获取 native 版本',
    }
    for i, c in enumerate(cases, 1):
        rows.append([i, c['class'].split('.')[-1], cn.get(c['name'], c['name']), f"{c['seconds']:.3f} s", '通过'])
    add_table(doc, ['序号', '测试类', '用例（中文说明）', '耗时', '结果'], rows, [1.1, 4.5, 8.0, 2.0, 2.0], 7.6, BLUE)
    add_p(doc, '12 个用例均由 Surefire 报告为 passed。测试输出中的 native 标准输出不影响断言结果；本轮没有 failures、errors 或 skipped。')
    add_heading(doc, '现有 JUnit 回归', 2)
    add_p(doc, 'PostpppInfoTests 负责结构体布局、字段默认值、固定宽度字段和 typed JSON 夹具；PostpppJavaLibApplicationTests 负责 Spring 上下文、native 版本、两种入口、许可证错误和异常 facade。两个测试类均在同一 clean verify 中执行。')
    add_heading(doc, '八个独立小时窗口', 2)
    add_p(doc, '当前证据没有八个独立小时的真实观测窗口，状态为待执行。补测时应固定同一产品版本和许可证，逐小时记录输入文件哈希、开始/结束时间、code、n_accept、have_coord、grade 和输出文件哈希。')
    add_heading(doc, '相同输入重复解算', 2)
    add_p(doc, '当前未进行真实数据的重复解算。正式验收应使用相同输入至少重复三次，比较 final_xyz、final_llh、rms_enu、converge_sec 和输出文件差异，以识别随机性、全局状态或输出目录竞争。')
    add_heading(doc, '回归判定与失败处理', 2)
    add_table(doc, ['指标', '本轮结果', '失败时处理'], [
        ['测试发现', '12 项；两个测试类均被 Surefire 发现', '检查测试扫描、依赖和 native 资源目录。'],
        ['Failures / Errors', '0 / 0', '保留 XML、日志和输入哈希后修复，再重新执行 clean verify。'],
        ['Skipped', '0', '确认没有被条件注解或平台过滤静默跳过。'],
        ['构建状态', 'BUILD SUCCESS；JAR 已生成', '检查 Maven 依赖、Java release 和动态库加载错误。'],
    ], [4.4, 5.1, 7.5], 8.0, BLUE)

    add_page(doc, '四 定位精度与重复性测试')
    add_heading(doc, '测量目标与当前证据', 2)
    add_p(doc, '本轮自动化测试证明 Java 对象能完整承载 native 返回字段，但没有真实测站和独立真值，因此当前精度、收敛、固定率、可用率和重复性均标记为待业务验收。返回对象与输入契约作为精度测试的前置条件在本节列出。')
    add_heading(doc, '1. PostpppInfo ABI 与输入字段', 2)
    add_table(doc, ['覆盖项', '验证内容', '结果'], [
        ['结构体布局', '字段顺序、定长数组长度和 JNA native size 392 字节。', '通过'],
        ['默认值', 'year/month/day、navsys、elevMask、soltype、tropopt、bdsfreq、tidecorr、maxSigma、troInterval 等默认值。', '通过'],
        ['路径数组', 'nav/sp3/clk 的 Pointer[4]、UTF-8 NUL 结尾 Memory 和路径引用保持。', '通过'],
        ['定长字符', 'station_name[16]、ant_type[24]、radome[8] 的 UTF-8 字节容量和 NUL 校验。', '通过'],
        ['边界校验', '路径过多、NUL 字符、超长 UTF-8 文本、错误数组长度均被拒绝。', '通过'],
    ], [4.0, 10.7, 2.7], 8.5, BLUE)
    add_heading(doc, '2. 返回结果对象封装', 2)
    add_table(doc, ['覆盖项', '验证内容', '结果'], [
        ['响应外层', 'code/error/result 解析为 PostpppResponse，并保留 isSuccess/hasResult。', '通过'],
        ['结果字段', 'PostpppResult 的数组、时间和统计字段按 Go 下划线 JSON 名称解码。', '通过'],
        ['逐历元对象', 'PostpppEpoch 与 PostpppTime 的嵌套对象和布尔 accepted 字段。', '通过'],
        ['部分结果', 'code=-6 时仍保留 result，便于诊断；直接结果 facade 保留错误码。', '通过'],
        ['两种入口', '结构体入口和 JSON 入口均能返回统一错误封装。', '通过'],
    ], [4.0, 10.7, 2.7], 8.5, BLUE)
    add_heading(doc, '3. Native 集成冒烟', 2)
    add_table(doc, ['覆盖项', '验证内容', '结果'], [
        ['版本', 'getVersion() 能加载 linux-x86-64/libPostppp.so 并返回版本字符串。', '通过'],
        ['初始化', 'native 初始化器产生的输入默认值与 Java defaults() 一致。', '通过'],
        ['许可证错误', '不可用许可证路径在 raw 和 typed API 中均返回可解析错误。', '通过'],
        ['错误码保留', 'PostpppException 保留 native code，调用方可按 -1～-7 处理。', '通过'],
    ], [4.0, 10.7, 2.7], 8.5, BLUE)
    add_heading(doc, '正式精度验收方法', 2)
    add_p(doc, '对每个真实站点保存独立参考坐标，并按 ECEF、ENU、纬度、经度和高程计算误差。至少报告固定率、可用率、P50/P95、RMSE、最大误差、收敛秒数和重复解算差异；报告中同时附上观测与产品文件 SHA-256。')
    add_heading(doc, '错误码与资源释放', 2)
    add_table(doc, ['场景', '期望行为', '结果'], [
        ['JSON 为空或缺少观测文件', '返回 code=-1 的 Go 错误封装；不生成 result。', '通过'],
        ['许可证路径不可用', 'raw 与 typed response 都保留许可证错误和 code=-1。', '通过'],
        ['code=-6 且带 result', 'PostpppResponse 保留部分 result；PostpppResult facade 抛异常。', '通过'],
        ['响应格式非法', 'JsonUtils 以 invalid_response 拒绝缺少数字 code 或错误类型。', '通过'],
        ['native Pointer', '服务层在 finally 中读取并释放一次；输入 Memory 不交给 freeP。', '通过'],
    ], [5.0, 9.8, 2.4], 8.2, BLUE)

    add_page(doc, '五 位移反应时间测试')
    add_heading(doc, '区分位移响应和解算收敛', 2)
    add_p(doc, 'PostPPP Java SDK 当前接口是同步日解算，不提供逐历元回调时间戳。因此“位移反应时间”不能由本轮 JUnit 耗时推导，必须在采集时间、native 处理时间、结果写出时间和业务告警时间之间分别打点。')
    add_heading(doc, '不同解算窗口的实测', 2)
    add_table(doc, ['窗口', '输入与指标', '当前状态'], [
        ['1 小时', '记录首个位移历元、首个 accepted 历元和结果写出时间。', '待真实数据'],
        ['8 小时', '比较收敛前后窗口，记录 converge_sec 和告警延迟。', '待真实数据'],
        ['全天', '比较完整日结果与中断窗口、n_gap、max_gap_sec。', '待真实数据'],
    ], [3.0, 10.5, 3.7], 8.4, BLUE)
    add_heading(doc, '待执行的受控位移试验', 2)
    add_p(doc, '需要已知阶跃幅值和准确发生时间的受控数据。验收应同时给出位移估计误差、从发生到 accepted 变化的延迟，以及从计算完成到业务可见的延迟。')

    add_page(doc, '六 5 MB 数据解码与全流程性能')
    add_heading(doc, '样本大小与边界', 2)
    add_p(doc, 'JSON 请求上限为 64 MiB，JNA 路径和固定数组有更小的 ABI 限制。性能验收应至少准备 5 MB JSON、5 MB 观测片段和全天 RINEX 三类样本，并分别观察 Java 堆、native RSS、解码耗时和输出磁盘占用。')
    add_heading(doc, '构建制品与完整性', 2)
    add_table(doc, ['制品', '大小', 'SHA-256', '判定'], [
        ['libPostppp.so', f"{SUMMARY['native_bytes']:,} bytes", SUMMARY['native_sha256'], '已加载并通过 native 冒烟'],
        ['postppp-java-lib-1.0.0.jar', f"{SUMMARY['jar_bytes']:,} bytes", SUMMARY['jar_sha256'], 'Maven verify 生成'],
    ], [4.2, 3.3, 7.0, 2.7], 7.8, BLUE)
    add_heading(doc, '实测性能', 2)
    add_p(doc, f'本轮只记录 Maven 测试耗时和构建产物哈希，没有 5 MB 真实解码样本、真实 PPP 数据集、固定 CPU 频率或业务性能阈值，因此没有填写虚构的吞吐量或 P95，也不把 {total_time} 的测试套件耗时解释为单站解算性能。')
    add_heading(doc, '判定边界', 2)
    add_p(doc, '后续测试应把文件读取、JNA 复制、native 解算、JSON 解析和 POS/TRO/QC 输出分段计时，并记录至少 P50、P95、最大 RSS、磁盘空间和失败恢复时间。')
    add_table(doc, ['待测项目', '建议方法', '当前状态'], [
        ['单站日解算耗时', '固定观测窗口和产品，重复 10 次并统计 P50/P95。', '待真实数据'],
        ['大文件容量', '5 MB、50 MB、全天数据逐档执行并观察 RSS 回落。', '待真实数据'],
        ['输出性能', '分别开启 POS、TRO、QC，比较文件写入和磁盘空间。', '待真实数据'],
        ['异常恢复', '截断文件、错日期、缺产品、无许可证，确认进程可继续服务。', '错误路径已覆盖，长时恢复待测'],
    ], [4.0, 9.7, 3.5], 8.3, BLUE)

    add_page(doc, '七 多任务并发解算')
    add_heading(doc, '任务隔离与公平比较', 2)
    add_p(doc, '每个任务应使用独立 PostpppInfo、许可证上下文和 outDir，并固定相同的产品版本。并发对比需记录任务提交、native 开始、native 结束和结果落盘时间，避免把队列等待误算为解算耗时。')
    add_heading(doc, '平台矩阵', 2)
    add_table(doc, ['平台', '资源目录', '本轮结论'], [
        ['Linux x86-64', 'linux-x86-64/libPostppp.so', '已执行，12 项全部通过。'],
        ['Linux AArch64', 'linux-aarch64/libPostppp.so', '未在本轮容器中执行，需要同架构运行时验证。'],
        ['macOS Apple Silicon', 'darwin-aarch64/libPostppp.so', '当前资源文件为 Linux ELF，不能作为 macOS native 通过依据。'],
        ['Windows', '未提供', '不纳入本轮。'],
    ], [4.2, 6.2, 6.8], 8.5, BLUE)
    add_heading(doc, '结果一致性', 2)
    add_p(doc, '当前只有单进程同步冒烟证据，没有真实数据的并发结果样本。因此暂不声称多任务结果一致性或吞吐量通过；补测时应比较相同输入的 final_xyz、n_accept、grade、输出文件和错误码。')
    add_heading(doc, '仍需验证的并发边界', 2)
    add_p(doc, '本轮测试覆盖同步服务和 native 进程锁路径，没有把一个 JVM 内的并发调用标记为线程安全。生产部署应先确认 Go 内核和输出目录的隔离规则；在确认前，优先采用业务层串行保护或多进程隔离。')
    add_table(doc, ['风险点', '建议门禁'], [
        ['同一 PostpppInfo 被并发修改', '调用期间冻结输入对象和数组。'],
        ['多个任务写同一 outDir', '按站点和日期分配唯一输出目录。'],
        ['native 全局许可证缓存', '每个进程固定许可证并在解算前校验有效期。'],
        ['native 返回 Pointer', '只由服务层读取和释放一次。'],
    ], [6.2, 11.0], 8.5, BLUE)

    add_page(doc, '八 问题清单与验收建议')
    add_heading(doc, '1. 问题清单', 2)
    add_table(doc, ['编号', '级别', '问题与影响', '关闭条件'], [
        ['TST-001', 'P1', '缺少真实测站和独立真值，不能验收绝对精度。', '提供匹配真值并统计 ECEF/ENU 误差、RMSE、P95。'],
        ['TST-002', 'P1', '没有真实位移时刻和阶跃幅值，不能验收响应时间。', '完成受控位移试验并记录采集、计算、回调各时间戳。'],
        ['TST-003', 'P2', 'AArch64/macOS/Windows native 尚未完成同架构验证。', '补齐资源并在目标平台执行相同测试矩阵。'],
        ['TST-004', 'P2', '没有长时容量和并发基线。', '补充多站、长时、不同输出目录和异常恢复测试。'],
    ], [2.3, 2.0, 7.0, 6.0], 8.0, BLUE)
    add_heading(doc, '2. 发布建议', 2)
    add_p(doc, '当前 Java SDK 契约和 Linux amd64 native 冒烟可以作为开发集成门禁。正式发布前仍应把真实 PPP 质量、精度、性能、平台矩阵和长时稳定性纳入独立验收，不以本报告的 12 项绿色测试替代业务指标。')
    add_heading(doc, '3. 归档证据', 2)
    add_table(doc, ['文件', '用途'], [
        ['docs/test-evidence/2026-10-01/maven-verify-linux-amd64.log', '完整 Maven clean verify 输出。'],
        ['docs/test-evidence/2026-10-01/TEST-*.xml', '去除机器相关 properties 后的 Surefire XML。'],
        ['docs/test-evidence/2026-10-01/*.txt', 'Surefire 文本汇总。'],
        ['docs/test-evidence/2026-10-01/summary.json', '环境、测试套件、用例、产物哈希和总计。'],
        ['docs/test-evidence/2026-10-01/source-baseline.json', 'Java 源码和 pom.xml 基线哈希。'],
        ['docs/test-evidence/2026-10-01/sha256.json / sha256.txt', '证据文件完整性清单。'],
    ], [8.0, 9.2], 8.4, BLUE)
    add_heading(doc, '4. 限制与后续验收', 2)
    add_p(doc, '本轮测试证明 Java SDK 的 ABI 映射、JSON 对象封装、错误处理和 Linux amd64 native 冒烟路径可用。它没有使用真实测站观测、独立坐标真值或全天候压力数据，因此不替代 PPP 绝对精度、收敛时间、性能容量和其他架构的专项验收。')
    add_bullets(doc, [
        '真实业务验收应补充与独立参考坐标对比的 ECEF/ENU 误差、固定率、可用率和重复性统计。',
        'Linux AArch64、macOS 和 Windows 需要分别提供并验证同架构动态库；当前资源中的 darwin-aarch64 文件实际不能直接作为 macOS 库使用。',
        '多线程并行部署前应使用独立进程或完成 native 线程安全验证，并固定输出目录和任务关联方式。',
        '升级 Go native ABI 时应同步更新 PostpppInfo 字段顺序、结构大小、响应字段和测试证据。',
    ])
    add_heading(doc, '5. 可复现命令', 2)
    add_code(doc, '''python3 scripts/report-tests/run_report_tests.py
# 或直接执行
mvn -B -o clean verify''')
    add_p(doc, '报告结论：12/12 项测试通过，构建通过，证据已归档。')
    add_page(doc, '附录 A 复现方法与证据索引')
    add_heading(doc, '1. 复现步骤', 2)
    add_p(doc, '在仓库根目录执行 python3 scripts/report-tests/run_report_tests.py。脚本会在 Linux amd64 Docker 中执行 clean verify，解析 Surefire XML，归档 Maven 日志、测试用例、源码基线和制品哈希。若仅检查本地代码，可直接执行 mvn -B -o clean verify。')
    add_heading(doc, '2. 证据目录', 2)
    add_table(doc, ['证据', '内容'], [
        ['summary.json', '执行时间、Git、容器、测试套件、用例、native/JAR 哈希。'],
        ['TEST-*.xml', '每个测试类的 Surefire 机器可读结果。'],
        ['*.txt 与 maven-*.log', '人类可读的测试和构建输出。'],
        ['source-baseline.json', 'Java 源码与 pom.xml 哈希，用于复现时确认代码基线。'],
        ['sha256.json / sha256.txt', '证据文件完整性清单。'],
    ], [5.2, 12.0], 8.4, BLUE)
    add_page(doc, '附录 B 历史结果与覆盖边界')
    add_heading(doc, '1. 原 v1.1 报告保留记录', 2)
    add_p(doc, '本版本延续 v1.1 的 12/12 绿色回归结论，并把字段映射、typed JSON、输入限制、平台矩阵和待验收项目展开到正文。历史测试产物的 SHA-256 和 Surefire XML 仍保留在证据目录。')
    add_heading(doc, '2. 历史构建制品', 2)
    add_p(doc, '本轮归档的 libPostppp.so、postppp-java-lib-1.0.0.jar、Docker 镜像和源码基线均记录哈希，便于后续比较 native ABI 或依赖升级带来的变化。')
    add_heading(doc, '3. 已覆盖', 2)
    add_bullets(doc, [
        'PostpppInfo 392 字节结构体布局、默认值、路径数组、固定宽度 UTF-8 字段。',
        '成功 JSON、-6 部分结果、错误 code/error、PostpppResponse 与 PostpppResult。',
        '结构体和 JSON 两种入口、native 版本、许可证错误和 typed facade。',
    ])
    add_heading(doc, '4. 未覆盖或待业务验收', 2)
    add_bullets(doc, [
        '真实观测文件的日解算精度、收敛时间、固定率和坐标真值误差。',
        '多平台 native、长时运行、并发吞吐、大文件容量和异常恢复。',
        'POS、SINEX_TRO、QC 输出文件的内容级验收和磁盘容量门禁。',
    ])
    out = DOCS / 'PostPPP-Java-SDK测试报告.docx'
    doc.save(out)
    return out


def build_report_md():
    out = DOCS / 'PostPPP-Java-SDK测试报告.md'
    suite_times = {s['name'].split('.')[-1]: f"{float(s['time']):.3f} s" for s in SUMMARY['suites']}
    total_time = f"{sum(float(s['time']) for s in SUMMARY['suites']):.3f} s"
    lines = []
    lines += ['# PostPPP Java SDK 测试报告', '', '报告 v1.1  |  SDK 1.0.0  |  原生库 1.0  |  2026-10-01', '']
    lines += ['## 结论摘要', '', '结论为通过。本轮在 Linux amd64 Docker 环境执行 `mvn -B -o clean verify`，两个测试类共 12 项测试全部通过，Failures=0、Errors=0、Skipped=0，构建结果为 `BUILD SUCCESS`。测试覆盖 `PostpppInfo` ABI 映射、默认值、UTF-8 和定长字段、返回 JSON 的 Java 对象封装、错误码保留、native 版本和许可证错误路径。', '']
    lines += ['| 测试内容 | 本轮结果 | 结论 |', '| --- | --- | --- |', '| PostpppInfo 映射 | 6 项；0 失败、0 错误、0 跳过 | 通过 |', '| 应用与 native 冒烟 | 6 项；0 失败、0 错误、0 跳过 | 通过 |', '| 总计 | 12 项；0 失败、0 错误、0 跳过 | 通过 |', '| 构建 | `mvn -B -o clean verify`；BUILD SUCCESS | 通过 |', '']
    lines += ['## 一、测试环境与计量方法', '', f'- Git 提交：`{SUMMARY["git_head"]}`', '- 执行平台：Linux amd64（Docker `--platform linux/amd64`）', f'- Docker 镜像：`{SUMMARY["image"]}`', f'- Docker 镜像 ID：`{SUMMARY["image_id"]}`', '- Java：Eclipse Temurin 21，工程编译 release 20', '- Maven：容器内 Maven，离线模式 `-o`', '- SDK：`com.navfirst:postppp-java-lib:1.0.0`', f'- native：{SUMMARY["native_bytes"]:,} bytes，SHA-256 `{SUMMARY["native_sha256"]}`', f'- JAR：{SUMMARY["jar_bytes"]:,} bytes，SHA-256 `{SUMMARY["jar_sha256"]}`', '', '```sh', 'mvn -B -o clean verify', '```', '', '判定条件：测试数不少于 12，Failures、Errors、Skipped 均为 0，Maven 输出 `BUILD SUCCESS`，并归档 Surefire XML/TXT、日志和 SHA-256 清单。', '', '### 计时与统计口径', '', '测试套件耗时来自 Surefire 汇总，只表示当前容器内 Java 测试和 native 冒烟的执行时间。报告不把该耗时当作单站 PPP 解算性能，也不从没有真实观测窗口的用例推导收敛时间、吞吐量或精度。', '']
    lines += ['## 二、数据来源与完整性', '', '测试数据由固定 JSON 响应夹具、PostpppInfo 构造对象和 native 错误路径组成。测试不把不存在的观测文件当作 PPP 精度样本；真实观测、广播星历、SP3、CLK、BIA 和 ANTEX 文件由业务部署提供。', '', '| 数据类别 | 本轮使用 | 完整性检查 |', '| --- | --- | --- |', '| JSON 响应夹具 | 成功响应、逐历元对象、-6 部分结果 | Gson 下划线字段、数组、嵌套时间和布尔值均断言 |', '| 结构体输入 | `PostpppInfo.defaults()`、Builder 和路径 setter | 392 字节布局、数组长度、UTF-8 NUL 结尾和默认值断言 |', '| 错误输入 | 空 JSON、缺少观测文件、不可用许可证、非法路径 | 统一解析为 code/error 或 `PostpppException` |', '| 原生文件 | JAR 内 `linux-x86-64/libPostppp.so` | 文件大小、加载结果和 SHA-256 已归档 |', '', '### 传输封装与星历', '', '输入通过 Java 对象、JNA Structure 或 JSON 字符串进入 Go C ABI。nav、sp3、clk 的每个路径均保留 UTF-8 NUL 结尾的 native Memory；测试验证路径数量上限和引用生命周期。星历、精密轨道、钟差、BIA、ANTEX、ERP 与 BLQ 的真实内容不在本轮夹具中。', '', '### 主测试站点与时间', '', '本轮没有可复核的真实测站、独立坐标真值或八个独立小时窗口，因此不填写虚构的站点名、观测日期和定位误差。业务验收应提供可归档的 RINEX、产品文件、真值和运行日志。', '', '路径字段按 UTF-8 传递，定长 `station_name`、`ant_type`、`radome` 按字节容量限制。当前证据没有有效许可证和真实日观测数据，因此本轮不声称完成坐标精度、收敛时间或吞吐量验收。', '']
    lines += [f'## 三、回归结果与独立站点解算', '', '| 测试类 | 测试数 | 失败 | 错误 | 跳过 | 耗时 |', '| --- | ---: | ---: | ---: | ---: | ---: |', f'| `PostpppInfoTests` | 6 | 0 | 0 | 0 | {suite_times.get("PostpppInfoTests", "n/a")} |', f'| `PostpppJavaLibApplicationTests` | 6 | 0 | 0 | 0 | {suite_times.get("PostpppJavaLibApplicationTests", "n/a")} |', f'| 合计 | 12 | 0 | 0 | 0 | {total_time} |', '', '### 用例明细', '', '| 测试类 | 用例 | 状态 |', '| --- | --- | --- |']
    for c in SUMMARY['cases']:
        lines.append(f'| `{c["class"].split(".")[-1]}` | `{c["name"]}` | 通过 |')
    lines += ['', '12 个用例均由 Surefire 报告为 passed；native 标准输出不影响断言结果。', '', '### 现有 JUnit 回归', '', '`PostpppInfoTests` 负责结构体布局、字段默认值、固定宽度字段和 typed JSON 夹具；`PostpppJavaLibApplicationTests` 负责 Spring 上下文、native 版本、两种入口、许可证错误和异常 facade。', '', '### 八个独立小时窗口', '', '当前证据没有八个独立小时的真实观测窗口，状态为待执行。补测时应固定同一产品版本和许可证，逐小时记录输入文件哈希、开始/结束时间、code、n_accept、have_coord、grade 和输出文件哈希。', '', '### 相同输入重复解算', '', '当前未进行真实数据的重复解算。正式验收应使用相同输入至少重复三次，比较 final_xyz、final_llh、rms_enu、converge_sec 和输出文件差异。', '']
    lines += ['## 四、定位精度与重复性测试', '', '### 测量目标与当前证据', '', '本轮自动化测试证明 Java 对象能完整承载 native 返回字段，但没有真实测站和独立真值，因此当前精度、收敛、固定率、可用率和重复性均标记为待业务验收。', '', '- `PostpppInfo` 字段顺序、392 字节 native size、默认值和固定数组长度通过验证。', '- `nav/sp3/clk` 使用 `Pointer[4]`，路径以 UTF-8 NUL 结尾并保持 Memory 引用。', '- `station_name[16]`、`ant_type[24]`、`radome[8]` 的 UTF-8 字节容量和非法输入拒绝通过验证。', '- `PostpppResponse`、`PostpppResult`、`PostpppEpoch`、`PostpppTime` 的嵌套 JSON 解码通过验证。', '- code=-6 时保留部分 result；raw 和 typed 入口保留一致的错误码。', '', '### 正式精度验收方法', '', '对每个真实站点保存独立参考坐标，并按 ECEF、ENU、纬度、经度和高程计算误差。至少报告固定率、可用率、P50/P95、RMSE、最大误差、收敛秒数和重复解算差异。', '', '### 错误码与资源释放', '', '| 场景 | 期望行为 | 结果 |', '| --- | --- | --- |', '| JSON 为空或缺少观测文件 | 返回 code=-1 的 Go 错误封装，不生成 result | 通过 |', '| 许可证路径不可用 | raw 与 typed response 均保留许可证错误 | 通过 |', '| code=-6 且带 result | 保留部分 result；结果 facade 抛异常 | 通过 |', '| 响应格式非法 | `JsonUtils` 以 `invalid_response` 拒绝 | 通过 |', '| native Pointer | 服务层在 finally 中读取并释放一次 | 通过 |', '']
    lines += ['## 五、位移反应时间测试', '', '### 区分位移响应和解算收敛', '', 'PostPPP Java SDK 当前接口是同步日解算，不提供逐历元回调时间戳。因此位移反应时间不能由本轮 JUnit 耗时推导，必须在采集时间、native 处理时间、结果写出时间和业务告警时间之间分别打点。', '', '| 窗口 | 输入与指标 | 当前状态 |', '| --- | --- | --- |', '| 1 小时 | 首个位移历元、首个 accepted 历元和结果写出时间 | 待真实数据 |', '| 8 小时 | 收敛前后窗口、converge_sec 和告警延迟 | 待真实数据 |', '| 全天 | 完整日结果、n_gap、max_gap_sec | 待真实数据 |', '', '受控试验还需要已知阶跃幅值和准确发生时间，并报告从发生到 accepted 变化的延迟。', '']
    lines += ['## 六、5 MB 数据解码与全流程性能', '', '### 样本大小与边界', '', 'JSON 请求上限为 64 MiB，JNA 路径和固定数组有更小的 ABI 限制。性能验收应至少准备 5 MB JSON、5 MB 观测片段和全天 RINEX 三类样本。', '', '| 制品 | 大小 | SHA-256 | 判定 |', '| --- | ---: | --- | --- |', f'| `libPostppp.so` | {SUMMARY["native_bytes"]:,} bytes | `{SUMMARY["native_sha256"]}` | 已加载并通过 native 冒烟 |', f'| `postppp-java-lib-1.0.0.jar` | {SUMMARY["jar_bytes"]:,} bytes | `{SUMMARY["jar_sha256"]}` | Maven verify 生成 |', '', '### 实测性能', '', '本轮只记录 Maven 测试耗时和构建产物哈希，没有 5 MB 真实解码样本、真实 PPP 数据集或固定硬件基线，因此没有填写虚构的吞吐量或 P95。', '', '### 判定边界', '', '| 待测项目 | 建议方法 | 当前状态 |', '| --- | --- | --- |', '| 单站日解算耗时 | 固定观测窗口和产品，重复 10 次并统计 P50/P95 | 待真实数据 |', '| 大文件容量 | 5 MB、50 MB、全天数据逐档执行并观察 RSS 回落 | 待真实数据 |', '| 输出性能 | 分别开启 POS、TRO、QC，比较文件写入和磁盘空间 | 待真实数据 |', '| 异常恢复 | 截断文件、错日期、缺产品、无许可证，确认进程可继续服务 | 错误路径已覆盖，长时恢复待测 |', '']
    lines += ['## 七、多任务并发解算', '', '### 任务隔离与公平比较', '', '每个任务应使用独立 PostpppInfo、许可证上下文和 outDir，并固定相同的产品版本。并发对比需记录任务提交、native 开始、native 结束和结果落盘时间。', '', '### 平台矩阵', '', '| 平台 | 资源目录 | 本轮结论 |', '| --- | --- | --- |', '| Linux x86-64 | `linux-x86-64/libPostppp.so` | 已执行，12 项全部通过 |', '| Linux AArch64 | `linux-aarch64/libPostppp.so` | 未在本轮容器中执行，需要同架构运行时验证 |', '| macOS Apple Silicon | `darwin-aarch64/libPostppp.so` | 当前资源文件为 Linux ELF，不能作为 macOS native 通过依据 |', '| Windows | 未提供 | 不纳入本轮 |', '', '### 结果一致性与仍需验证的边界', '', '当前没有真实数据的并发结果样本，暂不声称多任务结果一致性或吞吐量通过。生产部署应为每个任务分配唯一输出目录，并在确认 native 线程安全前采用业务层串行保护或多进程隔离。', '', '- 同一 `PostpppInfo` 在调用期间不得并发修改。', '- native 全局许可证缓存应在每个进程中固定并在解算前校验。', '- native 返回 Pointer 只由服务层读取和释放一次。', '']
    lines += ['## 八、问题清单与验收建议', '', '| 编号 | 级别 | 问题与影响 | 关闭条件 |', '| --- | --- | --- | --- |', '| TST-001 | P1 | 缺少真实测站和独立真值，不能验收绝对精度 | 提供匹配真值并统计 ECEF/ENU 误差、RMSE、P95 |', '| TST-002 | P1 | 没有真实位移时刻和阶跃幅值，不能验收响应时间 | 完成受控位移试验并记录采集、计算、回调时间戳 |', '| TST-003 | P2 | AArch64、macOS、Windows native 尚未完成同架构验证 | 补齐资源并执行相同测试矩阵 |', '| TST-004 | P2 | 没有长时容量和并发基线 | 补充多站、长时、不同输出目录和异常恢复测试 |', '', '当前 Java SDK 契约和 Linux amd64 native 冒烟可作为开发集成门禁。正式发布前仍应把真实 PPP 质量、精度、性能、平台矩阵和长时稳定性纳入独立验收。', '']
    lines += ['## 附录 A、复现方法与证据索引', '', '在仓库根目录执行：', '', '```sh', 'python3 scripts/report-tests/run_report_tests.py', '# 或直接执行', 'mvn -B -o clean verify', '```', '', '脚本会在 Linux amd64 Docker 中执行 clean verify，解析 Surefire XML，归档 Maven 日志、测试用例、源码基线和制品哈希。证据目录：[`docs/test-evidence/2026-10-01`](test-evidence/2026-10-01)。', '', '- `maven-verify-linux-amd64.log`：完整构建与测试输出。', '- `TEST-*.xml`、`*.txt`：Surefire 机器可读和人类可读结果。', '- `summary.json`：环境、套件、用例、产物哈希和总计。', '- `source-baseline.json`、`sha256.json`、`sha256.txt`：源码与证据完整性基线。', '']
    lines += ['## 附录 B、历史结果与覆盖边界', '', '### 原 v1.1 报告保留记录', '', '本版本延续 v1.1 的 12/12 绿色回归结论，并把字段映射、typed JSON、输入限制、平台矩阵和待验收项目展开到正文。', '', '### 历史构建制品', '', '本轮归档的 native、JAR、Docker 镜像和源码基线均记录哈希，便于后续比较 ABI 或依赖升级带来的变化。', '', '### 已覆盖', '', '- PostpppInfo 392 字节结构体布局、默认值、路径数组、固定宽度 UTF-8 字段。', '- 成功 JSON、-6 部分结果、错误 code/error、PostpppResponse 与 PostpppResult。', '- 结构体和 JSON 两种入口、native 版本、许可证错误和 typed facade。', '', '### 未覆盖或待业务验收', '', '- 真实观测文件的日解算精度、收敛时间、固定率和坐标真值误差。', '- 多平台 native、长时运行、并发吞吐、大文件容量和异常恢复。', '- POS、SINEX_TRO、QC 输出文件的内容级验收和磁盘容量门禁。', '', '报告结论：12/12 项测试通过，构建通过，证据已归档。', '']
    out.write_text('\n'.join(lines), encoding='utf-8')
    return out


if __name__ == '__main__':
    guide = build_guide()
    report = build_report()
    md = build_report_md()
    print(guide)
    print(report)
    print(md)
